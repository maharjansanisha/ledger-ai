"""Document indexing and search for Ask Your Ledger (bahikhata/rag.py).

FakeIndex stands in for Pinecone (no network); the manifest goes to tmp_path.
PDF and docx fixtures are built in memory, so no binary files are committed.
"""

import io

import pytest
from pinecone import PineconeError, UnauthorizedError

from bahikhata import config, rag
from tests.fakes import FakeIndex


@pytest.fixture(autouse=True)
def index(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "RAG_DIR", tmp_path / "rag")
    fake = FakeIndex(text_field="chunk_text")
    rag.set_index(fake, "chunk_text")
    yield fake
    rag.set_index(None)


def make_pdf(text: str) -> bytes:
    """A one-page PDF with `text` in Helvetica, with a correct xref table."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for n, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{n} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def make_blank_pdf() -> bytes:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def make_docx(*paragraphs: str) -> bytes:
    import docx

    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Rent", "Rs 40,000"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --- extract_text ------------------------------------------------------------

def test_extract_text_from_pdf():
    assert "Transport budget is Rs 5000" in rag.extract_text("budget.pdf", make_pdf("Transport budget is Rs 5000"))


def test_extract_text_from_docx_includes_tables():
    text = rag.extract_text("policy.docx", make_docx("Office supplies need approval."))
    assert "Office supplies need approval." in text
    assert "Rent | Rs 40,000" in text


@pytest.mark.parametrize("name", ["notes.txt", "notes.md", "prices.csv"])
def test_extract_text_from_plain_text_files(name):
    assert rag.extract_text(name, "चिया, Rs 30".encode()) == "चिया, Rs 30"


def test_scanned_pdf_without_text_is_no_text_error():
    with pytest.raises(rag.RagError) as exc:
        rag.extract_text("scan.pdf", make_blank_pdf())
    assert exc.value.kind == "no_text"


def test_corrupt_pdf_is_unreadable_error():
    with pytest.raises(rag.RagError) as exc:
        rag.extract_text("broken.pdf", b"not a pdf at all")
    assert exc.value.kind == "unreadable"


def test_unsupported_extension_is_rejected():
    with pytest.raises(rag.RagError):
        rag.extract_text("photo.jpg", b"\xff\xd8")


# --- chunk -------------------------------------------------------------------

def test_short_text_is_one_chunk():
    assert rag.chunk("One short paragraph.") == ["One short paragraph."]


def test_chunks_respect_size_overlap_and_word_boundaries():
    words = " ".join(f"word{i}" for i in range(2000))
    chunks = rag.chunk(words, size=500, overlap=100)
    assert len(chunks) > 1
    assert all(len(c) <= 500 for c in chunks)
    assert all(c.startswith("word") and not c.endswith("wor") for c in chunks)
    for first, second in zip(chunks, chunks[1:]):
        assert first.split()[-1] in second.split()  # consecutive chunks overlap


def test_chunks_prefer_paragraph_breaks():
    text = "A" * 300 + "\n\n" + "B" * 300
    assert rag.chunk(text, size=500, overlap=0) == ["A" * 300, "B" * 300]


# --- ingest / search / delete --------------------------------------------------

def test_ingest_indexes_chunks_and_records_manifest(index):
    result = rag.ingest("notes.txt", b"Transport budget is Rs 5000 per month.")
    assert (result.status, result.chunks) == ("indexed", 1)
    record = next(iter(index.records.values()))
    assert record["chunk_text"] == "Transport budget is Rs 5000 per month."
    assert record["doc_name"] == "notes.txt" and record["_id"].endswith("#0")
    [doc] = rag.list_documents()
    assert doc["name"] == "notes.txt" and doc["chunks"] == 1
    assert rag.has_documents()


def test_reingesting_the_same_file_is_a_duplicate_and_writes_nothing(index):
    rag.ingest("notes.txt", b"Transport budget is Rs 5000.")
    result = rag.ingest("copy-of-notes.txt", b"Transport budget is Rs 5000.")
    assert result.status == "duplicate"
    assert len(index.upsert_calls) == 1
    assert len(rag.list_documents()) == 1


def test_large_document_is_upserted_in_batches(index, monkeypatch):
    monkeypatch.setattr(config, "RAG_CHUNK_CHARS", 100)
    monkeypatch.setattr(config, "RAG_CHUNK_OVERLAP", 0)
    text = " ".join(f"word{i}" for i in range(3000)).encode()
    result = rag.ingest("big.txt", text)
    assert result.chunks > config.RAG_UPSERT_BATCH
    assert all(n <= config.RAG_UPSERT_BATCH for n in index.upsert_calls)
    assert sum(index.upsert_calls) == result.chunks


def test_ingest_failure_is_reported_not_raised_and_not_recorded(index):
    index.fail = PineconeError("boom")
    result = rag.ingest("notes.txt", b"some text")
    assert result.status == "failed"
    assert "unavailable" in result.message
    assert rag.list_documents() == []


def test_ingest_bad_api_key_says_so(index):
    index.fail = UnauthorizedError("401")
    result = rag.ingest("notes.txt", b"some text")
    assert "API key" in result.message


def test_ingest_blank_pdf_fails_with_no_text_message():
    result = rag.ingest("scan.pdf", make_blank_pdf())
    assert result.status == "failed" and "OCR" in result.message


def test_search_returns_relevant_chunks_above_min_score():
    rag.ingest("budget.txt", b"Transport budget is Rs 5000 per month.")
    rag.ingest("menu.txt", b"Momo plate costs Rs 180.")
    hits = rag.search("transport budget")
    assert [h.doc_name for h in hits] == ["budget.txt"]  # menu.txt scores 0, below RAG_MIN_SCORE
    assert hits[0].text.startswith("Transport budget")


def test_search_with_nothing_indexed_makes_no_call(index):
    index.fail = AssertionError("should not be called")
    assert rag.search("anything") == []


def test_search_failure_raises_rag_error(index):
    rag.ingest("budget.txt", b"Transport budget is Rs 5000.")
    index.fail = PineconeError("down")
    with pytest.raises(rag.RagError) as exc:
        rag.search("budget")
    assert exc.value.kind == "unavailable"


def test_delete_removes_vectors_and_manifest_entry(index):
    rag.ingest("budget.txt", b"Transport budget is Rs 5000.")
    rag.ingest("menu.txt", b"Momo plate costs Rs 180.")
    budget = next(d for d in rag.list_documents() if d["name"] == "budget.txt")
    rag.delete_document(budget["doc_id"])
    assert [d["name"] for d in rag.list_documents()] == ["menu.txt"]
    assert all(not record_id.startswith(budget["doc_id"]) for _ns, record_id in index.records)


def test_missing_config_is_config_error(monkeypatch):
    rag.set_index(None)
    monkeypatch.delenv("PINECONE_API_KEY", raising=False)
    with pytest.raises(rag.RagError) as exc:
        rag._connect()
    assert exc.value.kind == "config"

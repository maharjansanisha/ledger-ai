"""Document search for Ask Your Ledger: the ONLY module that talks to Pinecone.

Files attached in the Ask Your Ledger chat (pdf, docx, txt, md, csv) are turned
into text, split into overlapping chunks and upserted into an existing Pinecone
index that uses integrated embedding -- Pinecone embeds both the chunks and the
query, so this module never handles a vector. The name of the index's
semantic_text field is read from describe_index once per process.

Chunk ids are "<doc_id>#<n>", where doc_id is a prefix of the file's sha256, so
re-attaching the same file is a no-op and deleting a document needs no listing
call. The local manifest (data/rag/manifest.json) is the registry the page
lists documents from; Pinecone holds the text.

Documents are context only: ledger totals and figures always come from SQL
(ask.py), never from a retrieved chunk.

Errors: every Pinecone or parsing failure becomes RagError(kind, user_message),
mirroring llm_client.LLMError; the API key is never printed or logged.
"""

import hashlib
import io
import json
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import PurePath
from typing import Any

from bahikhata import config

SUPPORTED_EXTENSIONS = ("pdf", "docx", "txt", "md", "csv")


class RagError(Exception):
    """A document operation failed. `kind` is "config", "unavailable", "unreadable" or
    "no_text"; `user_message` is safe to show in the UI."""

    def __init__(self, kind: str, user_message: str, detail: str = ""):
        super().__init__(f"{kind}: {detail or user_message}")
        self.kind = kind
        self.user_message = user_message


_USER_MESSAGES = {
    "config": "Document search isn't configured — set PINECONE_API_KEY and PINECONE_INDEX in .env.",
    "auth": "Pinecone rejected the API key — check PINECONE_API_KEY.",
    "not_integrated": "The Pinecone index has no integrated embedding model — use an index created for a model.",
    "unavailable": "Document search is unavailable right now — try again shortly.",
    "unreadable": "I couldn't read that file.",
    "no_text": "That file has no extractable text (a scanned PDF needs OCR first).",
    "unsupported": "Unsupported file type.",
}


@dataclass
class Chunk:
    id: str
    doc_name: str
    text: str
    score: float


@dataclass
class IngestResult:
    name: str
    status: str              # "indexed", "duplicate" or "failed"
    chunks: int = 0
    message: str = ""        # user-safe; set when status is "failed"


# --- Pinecone connection (lazy; tests replace it with set_index) ------------

_connection: tuple[Any, str] | None = None


def set_index(index, text_field: str = "text") -> None:
    """Use `index` (e.g. tests.fakes.FakeIndex) instead of connecting; None resets."""
    global _connection
    _connection = None if index is None else (index, text_field)


def _connect() -> tuple[Any, str]:
    """(index, name of its semantic_text field). One describe call per process."""
    global _connection
    if _connection is not None:
        return _connection
    from pinecone import Pinecone

    try:
        api_key, name = config.get_pinecone_api_key(), config.get_pinecone_index()
    except RuntimeError as exc:
        raise RagError("config", _USER_MESSAGES["config"]) from exc
    with _pinecone_errors():
        pc = Pinecone(api_key=api_key)
        model = pc.describe_index(name)
        fields = model.schema.fields if model.schema else {}
        text_field = next((n for n, f in fields.items() if type(f).__name__ == "SemanticTextField"), None)
        if text_field is None:
            raise RagError("config", _USER_MESSAGES["not_integrated"], f"index {name!r}")
        _connection = (pc.Index(host=model.host), text_field)
    return _connection


class _pinecone_errors:
    """Context manager: any Pinecone exception -> RagError with a user-safe message."""

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc, tb):
        if exc is None or isinstance(exc, RagError):
            return False
        from pinecone import PineconeError, UnauthorizedError

        if isinstance(exc, UnauthorizedError):
            raise RagError("config", _USER_MESSAGES["auth"], type(exc).__name__) from exc
        if isinstance(exc, PineconeError):
            raise RagError("unavailable", _USER_MESSAGES["unavailable"], type(exc).__name__) from exc
        return False


# --- Manifest ----------------------------------------------------------------

def _manifest_path():
    return config.RAG_DIR / "manifest.json"


def _read_manifest() -> dict[str, dict[str, Any]]:
    path = _manifest_path()
    if not path.is_file():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def _write_manifest(manifest: dict[str, dict[str, Any]]) -> None:
    config.RAG_DIR.mkdir(parents=True, exist_ok=True)
    _manifest_path().write_text(json.dumps(manifest, ensure_ascii=False, indent=1), encoding="utf-8")


def list_documents() -> list[dict[str, Any]]:
    """Indexed documents, newest first: [{doc_id, name, chunks, size, added_at}, ...]."""
    docs = [{"doc_id": doc_id, **entry} for doc_id, entry in _read_manifest().items()]
    return sorted(docs, key=lambda d: d["added_at"], reverse=True)


def has_documents() -> bool:
    return bool(_read_manifest())


# --- Text extraction and chunking --------------------------------------------

def _extension(name: str) -> str:
    return PurePath(name).suffix.lower().lstrip(".")


def extract_text(name: str, data: bytes) -> str:
    """File bytes -> plain text, by extension. Raises RagError for unreadable/empty files."""
    ext = _extension(name)
    if ext not in SUPPORTED_EXTENSIONS:
        raise RagError("unreadable", _USER_MESSAGES["unsupported"], ext)
    try:
        if ext == "pdf":
            from pypdf import PdfReader

            text = "\n\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
        elif ext == "docx":
            import docx

            document = docx.Document(io.BytesIO(data))
            parts = [p.text for p in document.paragraphs]
            for table in document.tables:
                parts.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
            text = "\n\n".join(parts)
        else:
            text = data.decode("utf-8", errors="replace")
    except Exception as exc:  # pypdf / python-docx raise many unrelated types for a bad file
        raise RagError("unreadable", _USER_MESSAGES["unreadable"], type(exc).__name__) from exc
    if not text.strip():
        raise RagError("no_text", _USER_MESSAGES["no_text"])
    return text


def _normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def chunk(text: str, size: int | None = None, overlap: int | None = None) -> list[str]:
    """Split text into chunks of at most `size` chars, overlapping by about `overlap`.

    Each cut is moved back to the last paragraph, line, sentence or word break in
    the second half of the window, so chunks rarely end mid-word.
    """
    size = size or config.RAG_CHUNK_CHARS
    overlap = config.RAG_CHUNK_OVERLAP if overlap is None else overlap
    text = _normalize(text)
    chunks: list[str] = []
    start, n = 0, len(text)
    while start < n:
        end = min(start + size, n)
        if end < n:
            for sep in ("\n\n", "\n", ". ", " "):
                cut = text.rfind(sep, start + size // 2, end)
                if cut != -1:
                    end = cut + len(sep)
                    break
        piece = text[start:end].strip()
        if piece:
            chunks.append(piece)
        if end >= n:
            break
        next_start = max(end - overlap, start + 1)
        space = text.find(" ", next_start, end)
        start = space + 1 if space != -1 else next_start
    return chunks


# --- Ingest / search / delete ------------------------------------------------

def ingest(name: str, data: bytes) -> IngestResult:
    """Index one attached file. Never raises: failures come back as status "failed"."""
    doc_id = hashlib.sha256(data).hexdigest()[:16]
    manifest = _read_manifest()
    if doc_id in manifest:
        return IngestResult(name, "duplicate", manifest[doc_id]["chunks"])
    try:
        pieces = chunk(extract_text(name, data))
        if not pieces:
            raise RagError("no_text", _USER_MESSAGES["no_text"])
        index, text_field = _connect()
        records = [
            {"_id": f"{doc_id}#{n}", text_field: piece, "doc_id": doc_id, "doc_name": name, "chunk_no": n}
            for n, piece in enumerate(pieces)
        ]
        with _pinecone_errors():
            for i in range(0, len(records), config.RAG_UPSERT_BATCH):
                index.upsert_records(namespace=config.RAG_NAMESPACE, records=records[i:i + config.RAG_UPSERT_BATCH])
    except RagError as exc:
        return IngestResult(name, "failed", message=exc.user_message)
    manifest[doc_id] = {
        "name": name, "chunks": len(pieces), "size": len(data),
        "added_at": datetime.now(timezone.utc).isoformat(),
    }
    _write_manifest(manifest)
    return IngestResult(name, "indexed", len(pieces))


def search(question: str, top_k: int | None = None) -> list[Chunk]:
    """Top chunks for `question`, best first, dropping hits below RAG_MIN_SCORE.
    Returns [] without a network call when nothing is indexed. Raises RagError."""
    if not has_documents():
        return []
    index, text_field = _connect()
    with _pinecone_errors():
        response = index.search(
            namespace=config.RAG_NAMESPACE, top_k=top_k or config.RAG_TOP_K,
            inputs={"text": question}, fields=[text_field, "doc_name"],
        )
    chunks = []
    for hit in response.result.hits:
        fields = hit.fields or {}
        if hit.score >= config.RAG_MIN_SCORE and fields.get(text_field):
            chunks.append(Chunk(hit.id, fields.get("doc_name") or "?", fields[text_field], hit.score))
    return chunks


def delete_document(doc_id: str) -> None:
    """Remove a document's chunks from Pinecone, then from the manifest. Raises RagError."""
    manifest = _read_manifest()
    entry = manifest.get(doc_id)
    if entry is None:
        return
    index, _ = _connect()
    ids = [f"{doc_id}#{n}" for n in range(entry["chunks"])]
    with _pinecone_errors():
        for i in range(0, len(ids), 1000):
            index.delete(ids=ids[i:i + 1000], namespace=config.RAG_NAMESPACE)
    del manifest[doc_id]
    _write_manifest(manifest)
